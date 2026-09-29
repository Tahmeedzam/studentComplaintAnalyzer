import docx

doc = docx.Document('NLP_MINI_PROJECT_FINALL.docx')
print('Paragraphs count:', len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    print(f'P{i:02d} [Style: {p.style.name}]: "{p.text}"')
